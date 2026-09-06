import importlib.util
import io
import os
from pathlib import Path
import sys
import tempfile
import types
import unittest
from unittest.mock import Mock, patch

from openpyxl import load_workbook

SCRIPTS = Path(__file__).resolve().parents[1] / "scripts"
sys.path.insert(0, str(SCRIPTS))


class DepositExportTest(unittest.TestCase):
    def test_excel_deducts_collected_deposits_without_increasing_cost(self):
        self.check_export(False)

    def test_excel_includes_deposits_beyond_first_page(self):
        self.check_export(True)

    def check_export(self, full_page):
        records = {
            "export_jobs": [{"id": "job", "trip_id": 1, "target_id": "group"}],
            "trips": [{"id": 1, "title": "Trip"}],
            "expenses": [{"id": 206, "trip_id": 1, "payer_name": "A", "amount": 300,
                          "currency": "THB", "participants": ["A", "B", "C"],
                          "item_name": "Hotel", "tag": "#Hotel", "created_at": "2026-09-06T00:00:00Z"}],
            "expense_deposits": [
                {"expense_id": 206, "trip_id": 1, "payer_name": "A", "receiver_name": "A", "amount_minor": 10000, "currency": "THB"},
                {"expense_id": 206, "trip_id": 1, "payer_name": "B", "receiver_name": "A", "amount_minor": 5000, "currency": "THB"},
                {"expense_id": 206, "trip_id": 1, "payer_name": "C", "receiver_name": "A", "amount_minor": 2500, "currency": "THB"},
            ],
        }
        uploads = []
        if full_page:
            records["expense_deposits"][:0] = [dict(records["expense_deposits"][0], expense_id=1, amount_minor=0) for _ in range(1000)]

        # Fake only remote database/storage; generate and inspect the real workbook.
        class Query:
            def __init__(self, table): self.table, self.filters, self.bounds = table, {}, (0,999)
            def select(self, *_): return self
            def eq(self, key, value): self.filters[key] = value; return self
            def limit(self, *_): return self
            def order(self, *_, **__): return self
            def update(self, *_): return self
            def range(self, start, end): self.bounds = (start,end); return self
            def execute(self):
                rows = [r for r in records[self.table] if all(r.get(k) == v for k, v in self.filters.items())]
                return types.SimpleNamespace(data=rows[self.bounds[0]:self.bounds[1]+1])

        class Storage:
            def from_(self, _): return self
            def upload(self, **kwargs): uploads.append(kwargs["file"])
            def get_public_url(self, _): return "https://example.invalid/export.xlsx"

        def rpc(name, params):
            self.assertEqual(name, "get_expense_deposits")
            self.assertEqual(params, {"p_trip_id": 1, "p_expense_id": None})
            return types.SimpleNamespace(execute=lambda: types.SimpleNamespace(data=records["expense_deposits"]))

        client = types.SimpleNamespace(table=Query, storage=Storage(), rpc=rpc)
        transport = types.SimpleNamespace(get=Mock(side_effect=AssertionError("Unexpected FX request")), post=Mock(side_effect=AssertionError("Unexpected LINE push")))
        modules = patch.dict(sys.modules, {"supabase": types.SimpleNamespace(create_client=lambda *_: client), "requests": transport})
        modules.start()
        self.addCleanup(modules.stop)
        spec = importlib.util.spec_from_file_location("deposit_export_under_test", SCRIPTS / "export_trip_job.py")
        export = importlib.util.module_from_spec(spec)
        spec.loader.exec_module(export)
        with tempfile.TemporaryDirectory() as directory, patch.object(tempfile, "tempdir", directory), patch.dict(os.environ, {
            "EXPORT_JOB_ID": "job", "SUPABASE_URL": "https://example.invalid", "SUPABASE_KEY": "test", "LINE_CHANNEL_ACCESS_TOKEN": "",
        }):
            export.main()
        self.assertEqual(len(uploads), 1)
        workbook = load_workbook(io.BytesIO(uploads[0]), data_only=True)
        rows = [tuple(row[:3]) for row in workbook["รวมทุกวัน"].iter_rows(values_only=True)]
        self.assertIn(("B", "A", 50), rows)
        self.assertIn(("C", "A", 75), rows)
        self.assertTrue(any(row[0] == "A" and row[1] == 225 for row in rows))
        self.assertIn("เงินมัดจำ", workbook.sheetnames)
        expense_sheet = workbook[workbook.sheetnames[0]]
        self.assertEqual(expense_sheet.cell(2, 6).value, 300)


if __name__ == "__main__":
    unittest.main()
